import os
import sys
import yaml
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def convert_to_pdf(docx_path, pdf_path):
    if sys.platform == "win32":
        from docx2pdf import convert
        convert(docx_path, pdf_path)
    else:
        import subprocess
        out_dir = os.path.dirname(os.path.abspath(pdf_path))
        subprocess.run(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", out_dir, docx_path], check=True)

def add_bottom_border(paragraph, color_hex="1E293B", sz="8"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{sz}" w:space="2" w:color="{color_hex}"/></w:pBdr>')
    pPr.append(pBdr)

def load_source_of_truth():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    sot_dir = os.path.join(base_dir, "source_of_truth")
    with open(os.path.join(sot_dir, "profile.yaml"), "r", encoding="utf-8") as f:
        profile = yaml.safe_load(f)
    with open(os.path.join(sot_dir, "skills.yaml"), "r", encoding="utf-8") as f:
        skills = yaml.safe_load(f)
    with open(os.path.join(sot_dir, "bullets_bank.yaml"), "r", encoding="utf-8") as f:
        bullets_bank = yaml.safe_load(f)
    with open(os.path.join(sot_dir, "certs_projects.yaml"), "r", encoding="utf-8") as f:
        certs_proj = yaml.safe_load(f)
    bullets_by_id = {b["id"]: b for b in bullets_bank.get("bullets", [])}
    return profile, skills, bullets_by_id, certs_proj

PROFILES = {
    "spacious": {
        "name_pt": Pt(18.5), "contact_pt": Pt(8.8), "contact_after": Pt(6),
        "sec_hdr_pt": Pt(10.0), "sec_before": Pt(5.5), "sec_after": Pt(2.5),
        "summary_pt": Pt(9.0), "summary_spacing": 1.10, "summary_after": Pt(3.5),
        "cert_pt": Pt(8.8), "cert_id_pt": Pt(7.8), "cert_spacing": 1.05, "cert_before": Pt(1.5), "cert_after": Pt(1.5),
        "skills_pt": Pt(8.8), "skills_spacing": 1.06, "skills_after": Pt(1.8),
        "exp_role_pt": Pt(9.5), "exp_date_pt": Pt(8.8), "exp_sub_pt": Pt(8.8), "exp_sub_after": Pt(2.5),
        "exp_bullet_pt": Pt(8.8), "exp_bullet_spacing": 1.08, "exp_bullet_before": Pt(1.5), "exp_bullet_after": Pt(1.5),
        "proj_title_pt": Pt(9.2), "proj_tech_pt": Pt(8.0), "proj_bullet_pt": Pt(8.8), "proj_bullet_spacing": 1.08, "proj_bullet_before": Pt(1.5), "proj_bullet_after": Pt(1.5),
        "edu_deg_pt": Pt(8.8), "edu_date_pt": Pt(8.8), "edu_inst_pt": Pt(8.2), "edu_inst_after": Pt(2.5),
        "top_margin": Inches(0.40), "bottom_margin": Inches(0.35), "left_margin": Inches(0.50), "right_margin": Inches(0.50)
    },
    "compact": {
        "name_pt": Pt(17.5), "contact_pt": Pt(8.5), "contact_after": Pt(4),
        "sec_hdr_pt": Pt(9.5), "sec_before": Pt(4.5), "sec_after": Pt(2.0),
        "summary_pt": Pt(8.5), "summary_spacing": 1.04, "summary_after": Pt(2.5),
        "cert_pt": Pt(8.5), "cert_id_pt": Pt(7.5), "cert_spacing": 1.02, "cert_before": Pt(1.0), "cert_after": Pt(1.0),
        "skills_pt": Pt(8.5), "skills_spacing": 1.04, "skills_after": Pt(1.2),
        "exp_role_pt": Pt(9.0), "exp_date_pt": Pt(8.5), "exp_sub_pt": Pt(8.5), "exp_sub_after": Pt(2.0),
        "exp_bullet_pt": Pt(8.5), "exp_bullet_spacing": 1.04, "exp_bullet_before": Pt(1.0), "exp_bullet_after": Pt(1.0),
        "proj_title_pt": Pt(8.8), "proj_tech_pt": Pt(7.8), "proj_bullet_pt": Pt(8.5), "proj_bullet_spacing": 1.04, "proj_bullet_before": Pt(1.0), "proj_bullet_after": Pt(1.0),
        "edu_deg_pt": Pt(8.5), "edu_date_pt": Pt(8.5), "edu_inst_pt": Pt(8.0), "edu_inst_after": Pt(2.0),
        "top_margin": Inches(0.35), "bottom_margin": Inches(0.30), "left_margin": Inches(0.50), "right_margin": Inches(0.50)
    }
}

def render_document(config, out_docx, profile_name="spacious"):
    profile, _, bullets_by_id, certs_proj = load_source_of_truth()
    tp = PROFILES.get(profile_name, PROFILES["spacious"])
    
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = tp["top_margin"]
    section.bottom_margin = tp["bottom_margin"]
    section.left_margin = tp["left_margin"]
    section.right_margin = tp["right_margin"]

    primary_color = RGBColor(15, 23, 42)
    text_color = RGBColor(30, 41, 59)
    muted_color = RGBColor(71, 85, 105)

    # Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    p_name.paragraph_format.line_spacing = 1.0
    r_name = p_name.add_run(profile.get("name", "CANDIDATE NAME"))
    r_name.font.name = "Calibri"
    r_name.font.size = tp["name_pt"]
    r_name.font.bold = True
    r_name.font.color.rgb = primary_color

    # Contact Line
    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = tp["contact_after"]
    p_contact.paragraph_format.line_spacing = 1.0
    c = profile["contact"]
    contact_text = f"{c.get('location', '')}  |  {c.get('phone', '')}  |  {c.get('email', '')}  |  {c.get('linkedin', '')}"
    r_contact = p_contact.add_run(contact_text)
    r_contact.font.name = "Calibri"
    r_contact.font.size = tp["contact_pt"]
    r_contact.font.color.rgb = muted_color

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = tp["sec_before"]
        p.paragraph_format.space_after = tp["sec_after"]
        p.paragraph_format.line_spacing = 1.0
        add_bottom_border(p, color_hex="1E293B", sz="8")
        r = p.add_run(title)
        r.font.name = "Calibri"
        r.font.size = tp["sec_hdr_pt"]
        r.font.bold = True
        r.font.color.rgb = primary_color
        return p

    # 1. Summary
    add_section_header("PROFESSIONAL SUMMARY")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(1)
    p_sum.paragraph_format.space_after = tp["summary_after"]
    p_sum.paragraph_format.line_spacing = tp["summary_spacing"]
    
    summary_runs = config.get("summary_runs")
    if summary_runs:
        for item in summary_runs:
            r = p_sum.add_run(item.get("text", ""))
            r.font.name = "Calibri"
            r.font.size = tp["summary_pt"]
            r.font.bold = item.get("bold", False)
            r.font.color.rgb = text_color
    else:
        r = p_sum.add_run(config.get("summary", ""))
        r.font.name = "Calibri"
        r.font.size = tp["summary_pt"]
        r.font.color.rgb = text_color

    # 2. Certifications
    certs = certs_proj.get("certifications", [])
    if certs:
        add_section_header("CERTIFICATIONS")
        for cert in certs:
            p_c = doc.add_paragraph()
            p_c.paragraph_format.space_before = tp["cert_before"]
            p_c.paragraph_format.space_after = Pt(0)
            p_c.paragraph_format.line_spacing = tp["cert_spacing"]
            p_c.paragraph_format.left_indent = Inches(0.18)
            
            r_b = p_c.add_run("•  ")
            r_b.font.bold = True
            r_name = p_c.add_run(cert["name"])
            r_name.font.name = "Calibri"
            r_name.font.size = tp["cert_pt"]
            r_name.font.bold = True
            
            r_sep = p_c.add_run(" — ")
            r_iss = p_c.add_run(cert.get("issuer", ""))
            r_iss.font.name = "Calibri"
            r_iss.font.size = tp["cert_pt"]
            r_iss.font.italic = True
            
            cert_date = cert.get('issue_date') or cert.get('date', '')
            if cert_date:
                r_sp = p_c.add_run(f"    (Issued: {cert_date})")
                r_sp.font.name = "Calibri"
                r_sp.font.size = Pt(tp["cert_pt"].pt - 0.5)
                r_sp.font.color.rgb = muted_color

            cred_id = cert.get('credential_id') or cert.get('verification_id', '')
            if cred_id:
                p_id = doc.add_paragraph()
                p_id.paragraph_format.space_before = Pt(0)
                p_id.paragraph_format.space_after = tp["cert_after"]
                p_id.paragraph_format.line_spacing = 1.0
                p_id.paragraph_format.left_indent = Inches(0.32)
                r_id = p_id.add_run(f"ID: {cred_id}")
                r_id.font.name = "Consolas"
                r_id.font.size = tp["cert_id_pt"]
                r_id.font.color.rgb = muted_color

    # 3. Technical Skills
    add_section_header("TECHNICAL SKILLS")
    for cat, items in config.get("skills_lines", []):
        p_s = doc.add_paragraph()
        p_s.paragraph_format.space_before = Pt(0)
        p_s.paragraph_format.space_after = tp["skills_after"]
        p_s.paragraph_format.line_spacing = tp["skills_spacing"]
        
        r_cat = p_s.add_run(f"{cat}: ")
        r_cat.font.name = "Calibri"
        r_cat.font.size = tp["skills_pt"]
        r_cat.font.bold = True
        r_cat.font.color.rgb = primary_color
        
        r_it = p_s.add_run(items)
        r_it.font.name = "Calibri"
        r_it.font.size = tp["skills_pt"]
        r_it.font.color.rgb = text_color

    # 4. Work Experience
    add_section_header("WORK EXPERIENCE")
    emp_list = profile.get("employers", [{}])
    emp = emp_list[0] if emp_list else {}
    emp_company = emp.get("company", "TechCorp Solutions")
    emp_client = emp.get("client", "")
    emp_display = f"{emp_company} (Client: {emp_client})" if emp_client else emp_company
    emp_role = emp.get("titles", {}).get("de", emp.get("official_designation", "Software Engineer"))
    emp_dates = f"{emp.get('start_date', 'Jun 2023')} – {emp.get('end_date', 'Present')}"
    emp_loc = emp.get("location", "Bangalore, India")
    emp_proj = emp.get("project_de", "Enterprise Data Platform")

    p_exp_hdr = doc.add_paragraph()
    p_exp_hdr.paragraph_format.space_before = Pt(2)
    p_exp_hdr.paragraph_format.space_after = Pt(0)
    p_exp_hdr.paragraph_format.line_spacing = 1.0
    p_exp_hdr.paragraph_format.tab_stops.add_tab_stop(Inches(7.27), WD_TAB_ALIGNMENT.RIGHT)
    
    r_role = p_exp_hdr.add_run(emp_role)
    r_role.font.name = "Calibri"
    r_role.font.size = tp["exp_role_pt"]
    r_role.font.bold = True
    
    p_exp_hdr.add_run(" — ")
    r_comp = p_exp_hdr.add_run(emp_display)
    r_comp.font.name = "Calibri"
    r_comp.font.size = tp["exp_role_pt"]
    r_comp.font.italic = True
    
    p_exp_hdr.add_run("\t")
    r_dt = p_exp_hdr.add_run(emp_dates)
    r_dt.font.name = "Calibri"
    r_dt.font.size = tp["exp_date_pt"]
    r_dt.font.bold = True
    r_dt.font.color.rgb = primary_color

    p_proj_sub = doc.add_paragraph()
    p_proj_sub.paragraph_format.space_before = Pt(0)
    p_proj_sub.paragraph_format.space_after = tp["exp_sub_after"]
    p_proj_sub.paragraph_format.line_spacing = 1.0
    p_proj_sub.paragraph_format.tab_stops.add_tab_stop(Inches(7.27), WD_TAB_ALIGNMENT.RIGHT)
    
    r_p_sub = p_proj_sub.add_run(f"Project: {emp_proj}")
    r_p_sub.font.name = "Calibri"
    r_p_sub.font.size = tp["exp_sub_pt"]
    r_p_sub.font.bold = True
    
    p_proj_sub.add_run("\t")
    r_loc = p_proj_sub.add_run(emp_loc)
    r_loc.font.name = "Calibri"
    r_loc.font.size = Pt(tp["exp_sub_pt"].pt - 0.5)
    r_loc.font.italic = True
    r_loc.font.color.rgb = muted_color

    for b_info in config.get("experience_bullets", []):
        lead = b_info.get("lead", "")
        text = b_info.get("text", "")
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = tp["exp_bullet_before"]
        p_b.paragraph_format.space_after = tp["exp_bullet_after"]
        p_b.paragraph_format.line_spacing = tp["exp_bullet_spacing"]
        p_b.paragraph_format.left_indent = Inches(0.18)
        
        r_dot = p_b.add_run("•  ")
        r_dot.font.name = "Calibri"
        r_dot.font.size = tp["exp_bullet_pt"]
        if lead:
            r_ld = p_b.add_run(f"{lead}: ")
            r_ld.font.name = "Calibri"
            r_ld.font.size = tp["exp_bullet_pt"]
            r_ld.font.bold = True
        r_tx = p_b.add_run(text)
        r_tx.font.name = "Calibri"
        r_tx.font.size = tp["exp_bullet_pt"]
        r_tx.font.color.rgb = text_color

    # 5. Technical Projects
    add_section_header("TECHNICAL PROJECTS")
    p_pr_title = doc.add_paragraph()
    p_pr_title.paragraph_format.space_before = Pt(2)
    p_pr_title.paragraph_format.space_after = Pt(0)
    p_pr_title.paragraph_format.line_spacing = 1.0
    r_prt = p_pr_title.add_run("Scalable Data Ingestion & Analytics Pipeline")
    r_prt.font.name = "Calibri"
    r_prt.font.size = tp["proj_title_pt"]
    r_prt.font.bold = True

    p_pr_tech = doc.add_paragraph()
    p_pr_tech.paragraph_format.space_before = Pt(0)
    p_pr_tech.paragraph_format.space_after = Pt(2)
    p_pr_tech.paragraph_format.line_spacing = 1.0
    r_tech = p_pr_tech.add_run("Technologies: Python, PySpark, AWS S3, PostgreSQL, Docker")
    r_tech.font.name = "Calibri"
    r_tech.font.size = tp["proj_tech_pt"]
    r_tech.font.italic = True
    r_tech.font.color.rgb = muted_color

    for b_info in config.get("project_bullets", []):
        text = b_info.get("text", "") if isinstance(b_info, dict) else str(b_info)
        p_b = doc.add_paragraph()
        p_b.paragraph_format.space_before = tp["proj_bullet_before"]
        p_b.paragraph_format.space_after = tp["proj_bullet_after"]
        p_b.paragraph_format.line_spacing = tp["proj_bullet_spacing"]
        p_b.paragraph_format.left_indent = Inches(0.18)
        
        r_dot = p_b.add_run("•  ")
        r_dot.font.name = "Calibri"
        r_dot.font.size = tp["proj_bullet_pt"]
        r_tx = p_b.add_run(text)
        r_tx.font.name = "Calibri"
        r_tx.font.size = tp["proj_bullet_pt"]
        r_tx.font.color.rgb = text_color

    # 6. Education
    add_section_header("EDUCATION")
    for edu in profile.get("education", []):
        p_edu = doc.add_paragraph()
        p_edu.paragraph_format.space_before = Pt(1)
        p_edu.paragraph_format.space_after = Pt(0)
        p_edu.paragraph_format.line_spacing = 1.0
        p_edu.paragraph_format.tab_stops.add_tab_stop(Inches(7.27), WD_TAB_ALIGNMENT.RIGHT)
        r_deg = p_edu.add_run(edu.get("degree", ""))
        r_deg.font.name = "Calibri"
        r_deg.font.size = tp["edu_deg_pt"]
        r_deg.font.bold = True
        p_edu.add_run("\t")
        r_dts = p_edu.add_run(edu.get("dates", ""))
        r_dts.font.name = "Calibri"
        r_dts.font.size = tp["edu_date_pt"]
        r_dts.font.bold = True

        p_inst = doc.add_paragraph()
        p_inst.paragraph_format.space_before = Pt(0)
        p_inst.paragraph_format.space_after = tp["edu_inst_after"]
        p_inst.paragraph_format.line_spacing = 1.0
        inst_text = edu.get("institution", "")
        if edu.get("cgpa"):
            inst_text += f"  |  CGPA: {edu['cgpa']}"
        r_in = p_inst.add_run(inst_text)
        r_in.font.name = "Calibri"
        r_in.font.size = tp["edu_inst_pt"]
        r_in.font.color.rgb = muted_color

    os.makedirs(os.path.dirname(out_docx), exist_ok=True)
    doc.save(out_docx)

def render(config_path, output_dir):
    with open(config_path, "r", encoding="utf-8") as f:
        config = yaml.safe_load(f)
    os.makedirs(output_dir, exist_ok=True)
    out_docx = os.path.join(output_dir, "resume.docx")
    out_pdf = os.path.join(output_dir, "resume.pdf")

    # Render spacious first
    render_document(config, out_docx, profile_name="spacious")
    convert_to_pdf(out_docx, out_pdf)

    # Check 1-page compliance
    import pypdf
    reader = pypdf.PdfReader(out_pdf)
    if len(reader.pages) > 1:
        # Auto-calibrate to compact typography
        render_document(config, out_docx, profile_name="compact")
        convert_to_pdf(out_docx, out_pdf)
